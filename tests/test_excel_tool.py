import json
import subprocess
import sys
import tempfile
import types
import unittest
from pathlib import Path
from unittest import mock

import openpyxl

from excel_tool import (
    MAP_PATCHES_HEADERS,
    _mapping_patch_type,
    map_patches_workbook,
)


class ExcelMapPatchesTests(unittest.TestCase):
    def _git(self, repo, *args):
        result = subprocess.run(
            ["git", "-C", str(repo)] + list(args),
            check=True,
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            text=True,
        )
        return result.stdout.strip()

    def _init_repo(self, root, name, branch):
        repo = root / name
        repo.mkdir()
        subprocess.run(
            ["git", "init", str(repo)], check=True,
            stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True,
        )
        self._git(repo, "config", "user.name", "Test User")
        self._git(repo, "config", "user.email", "test@example.com")
        self._git(repo, "checkout", "-B", branch)
        return repo

    def _commit(self, repo, title, tag=None):
        count = len(list(repo.glob("commit-*.txt"))) + 1
        path = repo / "commit-{0}.txt".format(count)
        path.write_text("{0}\n{1}\n".format(title, count), encoding="utf-8")
        self._git(repo, "add", path.name)
        self._git(repo, "commit", "-m", title)
        commit_hash = self._git(repo, "rev-parse", "HEAD")
        if tag:
            self._git(repo, "tag", tag, commit_hash)
        return commit_hash

    def _save_source(self, path, rows):
        workbook = openpyxl.Workbook()
        worksheet = workbook.active
        worksheet.title = "Sheet1"
        headers = [
            "SR描述", "AR编号", "AR子编号", "标题描述", "版本模块",
            "Commit信息", "社区类型", "社区链接", "客户影响",
        ]
        worksheet.append(headers)
        for row in rows:
            worksheet.append([row.get(header, "") for header in headers])
        workbook.save(str(path))
        workbook.close()

    def _fixture(self, root):
        oe_repo = self._init_repo(root, "oe", "OLK-6.6")
        upstream_repo = self._init_repo(root, "upstream", "master")

        oe_semicolon = self._commit(
            oe_repo, "OE only feature; with semicolon", "6.6.0-1.0.0",
        )
        oe_missing = self._commit(
            oe_repo, "Missing upstream", "6.6.0-2.0.0",
        )
        oe_shared = self._commit(oe_repo, "Shared feature", "6.6.0-3.0.0")
        self._commit(oe_repo, "Duplicate feature", "6.6.0-4.0.0")
        oe_duplicate_new = self._commit(
            oe_repo, "Duplicate feature", "6.6.0-5.0.0",
        )

        upstream_semicolon = self._commit(
            upstream_repo, "OE only feature; with semicolon", "v6.1-rc1",
        )
        upstream_shared = self._commit(
            upstream_repo, "Shared feature", "v6.2-rc1",
        )
        upstream_duplicate = self._commit(
            upstream_repo, "Duplicate feature", "v6.3-rc1",
        )

        source_xlsx = root / "source.xlsx"
        self._save_source(source_xlsx, [
            {
                "SR描述": "SR one",
                "AR编号": "AR1",
                "AR子编号": "Feature_1",
                "标题描述": "AR one",
                "版本模块": "Mini-OS",
                "Commit信息": (
                    "{0};OE only\nfeature; with semicolon;2024-01-01;3\n"
                    "{1};Missing upstream;2024-01-02;4"
                ).format(oe_semicolon, oe_missing),
                "社区类型": "合入openEuler olk内核社区版本",
                "社区链接": "https://gitee.com/openeuler/kernel/pulls/123",
                "客户影响": "影响",
            },
            {
                "SR描述": "SR two",
                "AR编号": "AR2",
                "AR子编号": "Bug_2",
                "标题描述": "AR two",
                "版本模块": "Virtualization",
                "Commit信息": "{0};Shared feature;2024-01-03;5".format(
                    upstream_shared
                ),
                "社区类型": "合入上游linux内核社区版本",
            },
            {
                "SR描述": "SR three",
                "AR编号": "AR3",
                "AR子编号": "Clean up_3",
                "标题描述": "AR three",
                "版本模块": "Perf",
                "Commit信息": "{0};Duplicate feature;2024-01-04;6".format(
                    upstream_duplicate
                ),
                "社区类型": "合入上游linux内核社区版本",
            },
            {
                "SR描述": "SR duplicate",
                "AR编号": "AR4",
                "AR子编号": "Feature_4",
                "标题描述": "AR duplicate",
                "版本模块": "Mini-OS",
                "Commit信息": (
                    "{0};OE only feature; with semicolon;2024-01-05;7"
                ).format(oe_semicolon),
                "社区类型": "合入openEuler olk内核社区版本",
            },
            {
                "SR描述": "user row",
                "AR编号": "AR5",
                "AR子编号": "Feature_5",
                "标题描述": "skip me",
                "版本模块": "Qemu",
                "Commit信息": "abcdef123456;skip;2024-01-06;1",
                "社区类型": "合入Qemu社区版本",
            },
        ])
        return {
            "oe_repo": oe_repo,
            "upstream_repo": upstream_repo,
            "source": source_xlsx,
            "oe_semicolon": oe_semicolon,
            "oe_missing": oe_missing,
            "oe_shared": oe_shared,
            "oe_duplicate_new": oe_duplicate_new,
            "upstream_semicolon": upstream_semicolon,
            "upstream_shared": upstream_shared,
            "upstream_duplicate": upstream_duplicate,
        }

    def test_map_patches_expands_and_enriches_kernel_commits(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            fixture = self._fixture(root)
            source_xls = root / "source.xls"
            source_xls.write_bytes(fixture["source"].read_bytes())
            output = root / "mapped.xlsx"

            result = map_patches_workbook(
                source_xls,
                output,
                oe_repo=str(fixture["oe_repo"]),
                upstream_repo=str(fixture["upstream_repo"]),
            )

            workbook = openpyxl.load_workbook(str(output), data_only=False)
            worksheet = workbook["全量补丁"]
            values = list(worksheet.iter_rows(values_only=True))
            workbook.close()

        self.assertEqual(list(values[0]), MAP_PATCHES_HEADERS)
        self.assertEqual(len(values), 6)
        self.assertEqual(result["stats"]["source_rows_scanned"], 5)
        self.assertEqual(result["stats"]["source_rows_selected"], 4)
        self.assertEqual(result["stats"]["commit_records"], 5)
        self.assertEqual(result["stats"]["rows_written"], 5)

        first = values[1]
        self.assertEqual(first[0], "SR one")
        self.assertEqual(first[1], "AR one")
        self.assertEqual(first[3], "Feature")
        self.assertEqual(first[4:7], ("950", "内核态", "Mini-OS"))
        self.assertEqual(first[7], "OE only feature; with semicolon")
        self.assertEqual(first[8], fixture["oe_semicolon"][:12])
        self.assertEqual(first[10], "6.6.0-1.0.0")
        self.assertEqual(
            first[11], "https://gitcode.com/openeuler/kernel/pull/123"
        )
        self.assertEqual(first[12], fixture["upstream_semicolon"][:12])
        self.assertEqual(first[13], "v6.1-rc1")
        self.assertEqual(
            first[14], "AR1-Feature_1-{0}".format(fixture["oe_semicolon"][:12])
        )
        self.assertEqual(first[15:18], ("否", "影响", "未合入"))

        missing = values[2]
        self.assertEqual(missing[8], fixture["oe_missing"][:12])
        self.assertIsNone(missing[12])
        self.assertIsNone(missing[13])

        shared = values[3]
        self.assertEqual(shared[3], "Bug")
        self.assertEqual(shared[8], fixture["oe_shared"][:12])
        self.assertEqual(shared[12], fixture["upstream_shared"][:12])

        duplicate = values[4]
        self.assertEqual(duplicate[3], "Cleanup")
        self.assertEqual(duplicate[8], fixture["oe_duplicate_new"][:12])
        self.assertEqual(duplicate[12], fixture["upstream_duplicate"][:12])
        self.assertEqual(result["stats"]["ambiguous_oe"], 1)

        duplicate_ar = values[5]
        self.assertEqual(duplicate_ar[8], fixture["oe_semicolon"][:12])
        self.assertTrue(str(duplicate_ar[14]).startswith("AR4-Feature_4-"))
        self.assertTrue(
            any(issue["kind"] == "upstream_unmatched" for issue in result["issues"])
        )

    def test_output_overwrite_requires_force(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            fixture = self._fixture(root)
            output = root / "mapped.xlsx"
            output.write_bytes(b"existing")
            with self.assertRaisesRegex(RuntimeError, "--force"):
                map_patches_workbook(
                    fixture["source"], output,
                    oe_repo=str(fixture["oe_repo"]),
                    upstream_repo=str(fixture["upstream_repo"]),
                )

    def test_rejects_legacy_xls_and_missing_headers(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            fixture = self._fixture(root)
            legacy = root / "legacy.xls"
            legacy.write_bytes(b"\xd0\xcf\x11\xe0legacy")
            with self.assertRaisesRegex(RuntimeError, "BIFF"):
                map_patches_workbook(
                    legacy, root / "legacy-output.xlsx",
                    oe_repo=str(fixture["oe_repo"]),
                    upstream_repo=str(fixture["upstream_repo"]),
                )

            missing = root / "missing.xlsx"
            workbook = openpyxl.Workbook()
            workbook.active.append(["wrong"])
            workbook.save(str(missing))
            workbook.close()
            with self.assertRaisesRegex(RuntimeError, "映射所需表头"):
                map_patches_workbook(
                    missing, root / "missing-output.xlsx",
                    oe_repo=str(fixture["oe_repo"]),
                    upstream_repo=str(fixture["upstream_repo"]),
                )

            with self.assertRaisesRegex(RuntimeError, "同一个文件"):
                map_patches_workbook(
                    fixture["source"], fixture["source"],
                    oe_repo=str(fixture["oe_repo"]),
                    upstream_repo=str(fixture["upstream_repo"]),
                )

    def test_invalid_ref_fails_without_output(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            fixture = self._fixture(root)
            output = root / "mapped.xlsx"
            with self.assertRaisesRegex(RuntimeError, "无法读取提交仓库"):
                map_patches_workbook(
                    fixture["source"], output,
                    oe_repo=str(fixture["oe_repo"]),
                    upstream_repo=str(fixture["upstream_repo"]),
                    oe_ref="missing-ref",
                )
            self.assertFalse(output.exists())

    def test_patch_type_normalization(self):
        self.assertEqual(_mapping_patch_type("Feature_1"), "Feature")
        self.assertEqual(_mapping_patch_type("Bug_22"), "Bug")
        self.assertEqual(_mapping_patch_type("Clean up_3"), "Cleanup")

    def test_oe_title_falls_back_to_source_and_formula_text_stays_literal(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            oe_repo = self._init_repo(root, "oe", "OLK-6.6")
            self._commit(oe_repo, "Unrelated OE commit")
            upstream_repo = self._init_repo(root, "upstream", "master")
            upstream_hash = self._commit(upstream_repo, "Canonical upstream title")
            source = root / "source.xlsx"
            self._save_source(source, [{
                "SR描述": "=1+1",
                "AR编号": "AR6",
                "AR子编号": "Feature_6",
                "标题描述": "  keep surrounding spaces  ",
                "版本模块": "Kernel",
                "Commit信息": (
                    "{0};Source fallback title;2024-01-06;1"
                ).format(upstream_hash),
                "社区类型": "合入上游linux内核社区版本",
                "客户影响": "=HYPERLINK(\"https://invalid/\",\"x\")",
            }])
            output = root / "mapped.xlsx"
            result = map_patches_workbook(
                source, output, oe_repo=str(oe_repo),
                upstream_repo=str(upstream_repo),
            )

            workbook = openpyxl.load_workbook(str(output), data_only=False)
            worksheet = workbook["全量补丁"]
            row = list(worksheet.iter_rows(min_row=2, max_row=2))[0]
            values = [cell.value for cell in row]
            data_types = [cell.data_type for cell in row]
            workbook.close()

        self.assertEqual(values[0], "=1+1")
        self.assertEqual(values[1], "  keep surrounding spaces  ")
        self.assertEqual(values[7], "Source fallback title")
        self.assertIsNone(values[8])
        self.assertEqual(values[12], upstream_hash[:12])
        self.assertEqual(values[16], "=HYPERLINK(\"https://invalid/\",\"x\")")
        self.assertEqual(data_types[0], "s")
        self.assertEqual(data_types[16], "s")
        self.assertEqual(result["stats"]["oe_unmatched"], 1)

    def test_mcp_mapping_returns_core_result_shape(self):
        class FakeFastMCP:
            def __init__(self, *args, **kwargs):
                self.settings = types.SimpleNamespace()

            def tool(self):
                return lambda function: function

        mcp_module = types.ModuleType("mcp")
        server_module = types.ModuleType("mcp.server")
        fastmcp_module = types.ModuleType("mcp.server.fastmcp")
        fastmcp_module.FastMCP = FakeFastMCP
        expected = {
            "output_file": "/tmp/mapped.xlsx",
            "stats": {"commit_records": 2, "rows_written": 2},
            "issues": [{"kind": "oe_unmatched", "sheet": "Sheet1", "source_row": 2}],
        }
        modules = {
            "mcp": mcp_module,
            "mcp.server": server_module,
            "mcp.server.fastmcp": fastmcp_module,
        }
        with mock.patch.dict(sys.modules, modules):
            sys.modules.pop("mcp_server", None)
            import mcp_server

            with mock.patch.object(
                mcp_server, "map_patches_workbook", return_value=expected,
            ) as mapper:
                result = json.loads(mcp_server.excel_map_patches(
                    "source.xls", "/tmp/mapped.xlsx", "/repo/oe", "/repo/linux",
                ))
            sys.modules.pop("mcp_server", None)

        self.assertEqual(result, expected)
        mapper.assert_called_once_with(
            input_file=Path("source.xls"),
            output_file=Path("/tmp/mapped.xlsx"),
            oe_repo="/repo/oe",
            upstream_repo="/repo/linux",
            oe_ref="OLK-6.6",
            upstream_ref="master",
            product_version="950",
            force=False,
        )


if __name__ == "__main__":
    unittest.main()
